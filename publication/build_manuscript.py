"""Build a readable manuscript using the supplied Jinhan DOCX's typography."""

import argparse
import json
import re
import textwrap
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor
from docx.text.run import Run

from validation_table import (COLUMN_FRACTIONS, HEADERS, LABELS, METRICS, SUMMARY_NOTE,
                              TABLE_TITLES, best_means, dimension_groups,
                              dimension_label, group_key, partition_summary, read_summary)


CITATION = re.compile(r"(\[\d+(?:\s*[,–-]\s*\d+)*\])")
BODY_LINE_SPACING = 251 / 240
FONT = "Times New Roman"


def word_count(text):
    return sum(any(character.isalnum() for character in word) for word in text.split())


def submission_fields(content, figures, output):
    sections = [
        ("Introduction", content["introduction"]),
        ("Methods", [label + ". " + text for label, text in content["methods"]]),
        ("Results", content["results"]),
        ("Discussion and Conclusion", content["discussion"]),
    ]
    counts = {
        "title_characters": len(content["title"]),
        "body_words": sum(word_count(text) for _, texts in sections for text in texts),
        "body_words_with_headings": sum(word_count(title) + sum(word_count(text) for text in texts)
                                        for title, texts in sections),
        "synopsis_words_with_labels": sum(word_count(label + " " + text)
                                           for label, text in content["synopsis"]),
        "impact_words": word_count(content["impact"]),
    }
    if (counts["title_characters"] > 125 or counts["body_words_with_headings"] > 750
            or counts["synopsis_words_with_labels"] > 100 or counts["impact_words"] > 40):
        raise ValueError(f"ISMRM 2027 text limit exceeded: {counts}")
    captions = [TABLE_TITLES[kind] + " " + SUMMARY_NOTE for kind in ("direct", "residual")]
    captions += [figure["caption"] for figure in figures]
    if any(len(text) > 500 for text in captions):
        raise ValueError("ISMRM 2027 captions must not exceed 500 characters.")
    lines = ["TITLE", content["title"], "", "SYNOPSIS"]
    lines += [label + ": " + text for label, text in content["synopsis"]]
    lines += ["", "IMPACT", content["impact"], "", "ABSTRACT BODY"]
    for title, texts in sections:
        lines += ["", title]
        for text in texts:
            lines += [CITATION.sub(lambda match: "$^{" + match.group(1) + "}$", text), ""]
    lines += ["REFERENCES"] + [f"{i}. {text}" for i, text in enumerate(content["references"], 1)]
    if content.get("code_availability_url"):
        lines += ["", "ACKNOWLEDGEMENTS", "The code is available at " + content["code_availability_url"] + "."]
    lines += ["", "FIGURE AND TABLE CAPTIONS"] + captions
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return counts


def text_value(value):
    return value.replace("+/-", "±").replace(" x ", " × ")


def template_properties(path):
    # Only formatting XML is needed; reference images are neither loaded nor copied.
    with ZipFile(path) as package:
        root = parse_xml(package.read("word/document.xml"))
        styles = parse_xml(package.read("word/styles.xml"))
        settings = parse_xml(package.read("word/settings.xml"))
    body = root.find(qn("w:body"))
    paragraphs = body.findall(qn("w:p"))
    texts = ["".join(p.itertext(tag=qn("w:t"))).strip() for p in paragraphs]
    for heading in ("Introduction", "Methods", "Results", "Conclusion", "References"):
        if not any(re.fullmatch(r"(?:\d+\s*\|\s*)?" + heading, text) for text in texts):
            raise ValueError("Use the Jinhan SCORD reference DOCX for manuscript formatting.")
    title_run = paragraphs[0].find(qn("w:r"))
    if title_run is None or title_run.find(qn("w:rPr")) is None:
        raise ValueError("The reference must provide explicit title typography.")
    title_props = title_run.find(qn("w:rPr"))
    if title_props.find(qn("w:sz")).get(qn("w:val")) != "32":
        raise ValueError("Expected the supplied 16-point title style.")
    section = body.find(qn("w:sectPr"))
    defaults = styles.find(qn("w:docDefaults"))
    if section is None or defaults is None:
        raise ValueError("Reference page geometry and style defaults are required.")
    intro = next(i for i, text in enumerate(texts) if re.fullmatch(r"(?:\d+\s*\|\s*)?Introduction", text))
    refs = texts.index("References")
    profiles = {
        "title": paragraphs[0], "authors": paragraphs[1], "affiliation": paragraphs[2],
        "heading": paragraphs[intro], "body": paragraphs[intro + 1],
        "references_heading": paragraphs[refs], "reference": paragraphs[refs + 1],
        "caption": next(p for p, text in zip(paragraphs, texts) if text.startswith("Figure 1.")),
    }
    normal = next(s for s in styles.findall(qn("w:style"))
                  if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) == "1")
    return section, defaults, profiles, settings, normal


def apply_template_typography(doc, profiles, content):
    affiliations = {number + value for number, value in content.get("affiliations", {}).items()}
    for index, p in enumerate(doc.paragraphs):
        value = p.text.strip()
        if not value:
            continue
        if index < 2:
            role = ("title", "authors")[index]
        elif value in affiliations:
            role = "affiliation"
        elif value in ("References", "Acknowledgements"):
            role = "references_heading"
        elif re.match(r"^\d+\s*\|", value):
            role = "heading"
        elif re.match(r"^\[\d+\]", value):
            role = "reference"
        elif re.match(r"^(?:Figure|Table)\s+\d+\.", value):
            role = "caption"
        else:
            role = "body"
        source = profiles[role]
        props = deepcopy(source.find(qn("w:pPr")))
        style = props.find(qn("w:pStyle"))
        if style is not None:
            props.remove(style)
        old = p._p.find(qn("w:pPr"))
        if old is not None:
            for tag in ("keepNext", "keepLines", "pageBreakBefore", "widowControl"):
                for existing in list(props.findall(qn("w:" + tag))):
                    props.remove(existing)
                control = old.find(qn("w:" + tag))
                if control is not None:
                    props.append(deepcopy(control))
            p._p.remove(old)
        p._p.insert(0, props)
        source_run = next(r for r in source.iter(qn("w:r")) if r.find(qn("w:t")) is not None)
        source_props = source_run.find(qn("w:rPr"))
        for run in p._p.iter(qn("w:r")):
            target_props = run.get_or_add_rPr()
            for tag in ("rFonts", "sz", "szCs"):
                current = target_props.find(qn("w:" + tag))
                if current is not None:
                    target_props.remove(current)
                original = source_props.find(qn("w:" + tag))
                if original is not None:
                    target_props.append(deepcopy(original))


def set_font(run, size=12, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._r.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + attr), FONT)
    for tag, value in (("szCs", str(int(size * 2))), ("bCs", "1" if bold else "0")):
        element = props.find(qn("w:" + tag))
        if element is None:
            element = OxmlElement("w:" + tag)
            props.append(element)
        element.set(qn("w:val"), value)
    return run


def add_text(p, text, bold=False, size=12, citations=True):
    chunks = CITATION.split(text_value(text)) if citations else (text_value(text),)
    for index, chunk in enumerate(chunks):
        if not chunk:
            continue
        if citations and index + 1 < len(chunks) and CITATION.fullmatch(chunks[index + 1]):
            chunk = chunk.rstrip()
        is_citation = citations and CITATION.fullmatch(chunk)
        run = set_font(p.add_run(chunk), size, bold)
        if is_citation:
            run.font.superscript = True


def add_hyperlink(p, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    element = OxmlElement("w:r")
    link.append(element)
    run = set_font(Run(element, p))
    run.text = url
    run.font.underline = True
    p._p.append(link)


def paragraph(doc, text="", label=None, size=12, bold=False, citations=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = BODY_LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.widow_control = True
    if label:
        add_text(p, label + ". ", bold=True, size=size, citations=False)
    add_text(p, text, bold=bold, size=size, citations=citations)
    return p


def heading(doc, text, new_page=False):
    p = paragraph(doc, text, size=14, bold=True, citations=False)
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = 1.0
    pf.space_before = Pt(12)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    pf.keep_together = True
    pf.page_break_before = new_page
    return p


def caption(doc, text, page_break=False):
    p = paragraph(doc)
    p.paragraph_format.keep_together = True
    p.paragraph_format.page_break_before = page_break
    lead = re.match(r"^((?:Figure|Table)\s+\d+\.\s+.*?\.)(?:\s+|$)", text)
    if lead:
        add_text(p, lead.group(1), bold=True, citations=False)
        add_text(p, text[len(lead.group(1)):], citations=False)
    else:
        add_text(p, text, citations=False)
    return p


def border(parent, side, value="single", size=8):
    element = OxmlElement("w:" + side)
    element.set(qn("w:val"), value)
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:color"), "000000")
    parent.append(element)


def add_table(doc, kind, rows, new_page=False):
    title = caption(doc, TABLE_TITLES[kind] + " " + SUMMARY_NOTE, page_break=new_page)
    title.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    section = doc.sections[-1]
    usable = section.page_width - section.left_margin - section.right_margin
    for column, fraction in zip(table.columns, COLUMN_FRACTIONS):
        column.width = int(usable * fraction)
    table_borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border(table_borders, side, "single" if side in ("top", "bottom") else "nil", 10)
    table._tbl.tblPr.append(table_borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        margin = OxmlElement("w:" + side)
        margin.set(qn("w:w"), "40")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    table._tbl.tblPr.append(margins)
    for cell, text in zip(table.rows[0].cells, HEADERS):
        cell.text = text
        borders = OxmlElement("w:tcBorders")
        border(borders, "bottom", size=6)
        cell._tc.get_or_add_tcPr().append(borders)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    best = best_means(rows)
    for row in rows:
        cells = table.add_row().cells
        cells[0].text = LABELS[row["method"]]
        cells[1].text = dimension_label(row)
        for cell, (metric, digits, _) in zip(cells[2:], METRICS):
            mean, std = float(row[metric + "_mean"]), float(row[metric + "_std"])
            p = cell.paragraphs[0]
            set_font(p.add_run(f"{mean:.{digits}f}"), 10.5, mean == best[metric])
            set_font(p.add_run(f"\n± {std:.{digits}f}"), 10.5, False)
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column, cell in enumerate(row.cells):
            cell.width = int(usable * COLUMN_FRACTIONS[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.space_before = Pt(4)
                pf.space_after = Pt(4)
                pf.line_spacing = 1.0
                pf.keep_with_next = index < len(table.rows) - 1
                for run in p.runs:
                    set_font(run, 10.5, True if index == 0 else bool(run.bold))
    for members in dimension_groups(rows):
        if len(members) == 1:
            continue
        first, last = members[0][0] + 1, members[-1][0] + 1
        label = dimension_label(members[0][1])
        cell = table.cell(first, 1).merge(table.cell(last, 1))
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            set_font(run, 10.5)
    for index in range(1, len(rows)):
        major = group_key(rows[index - 1]) != group_key(rows[index])
        for column in range(len(HEADERS)) if major else (0, *range(2, len(HEADERS))):
            props = table.rows[index + 1]._tr.tc_lst[column].get_or_add_tcPr()
            borders = props.find(qn("w:tcBorders"))
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                props.append(borders)
            border(borders, "top", size=6 if major else 3)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                props = p._p.get_or_add_pPr()
                mark = props.find(qn("w:rPr"))
                if mark is None:
                    mark = OxmlElement("w:rPr")
                    props.append(mark)
                for tag in ("sz", "szCs"):
                    element = mark.find(qn("w:" + tag))
                    if element is None:
                        element = OxmlElement("w:" + tag)
                        mark.append(element)
                    element.set(qn("w:val"), "21")
    return table


def add_figure(doc, figure, figures_dir):
    section = doc.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    usable_height = section.page_height - section.top_margin - section.bottom_margin
    image_file = figures_dir / figure["file"]
    source = Image.from_file(str(image_file))
    ratio = source.px_width / source.px_height
    caption_lines = max(1, len(textwrap.wrap(figure["caption"], width=max(1, int(usable_width / Pt(5.2))))))
    grid = section._sectPr.find(qn("w:docGrid"))
    grid_line = int(grid.get(qn("w:linePitch"), "0")) / 20 if grid is not None else 0
    line_height = max(12 * BODY_LINE_SPACING, grid_line)
    available_height = usable_height - Pt(caption_lines * line_height + 32)
    if available_height <= Pt(72):
        raise ValueError("The figure caption leaves insufficient image space on the page.")
    width = min(usable_width, available_height * ratio)
    p = paragraph(doc)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.page_break_before = True
    pf.keep_with_next = True
    pf.keep_together = True
    pf.line_spacing = 1.0
    pf.space_after = Pt(0)
    p.add_run().add_picture(str(image_file), width=int(width))
    caption(doc, figure["caption"])


def build(args):
    content = json.loads(args.content.read_text(encoding="utf-8"))
    tables = partition_summary(read_summary(args.summary))
    figures = json.loads(args.figures_manifest.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    counts = submission_fields(content, figures,
                               args.output.with_name(args.output.stem + "_ISMRM2027_Fields.txt"))
    template_section, defaults, profiles, settings, template_normal = template_properties(args.template)
    doc = Document()
    old_defaults = doc.styles.element.find(qn("w:docDefaults"))
    if old_defaults is not None:
        doc.styles.element.remove(old_defaults)
    doc.styles.element.insert(0, deepcopy(defaults))
    normal = doc.styles["Normal"]
    for tag in ("pPr", "rPr"):
        old = normal.element.find(qn("w:" + tag))
        if old is not None:
            normal.element.remove(old)
        original = template_normal.find(qn("w:" + tag))
        if original is not None:
            normal.element.append(deepcopy(original))
    for tag in ("zoom", "defaultTabStop", "characterSpacingControl", "compat", "themeFontLang"):
        old = doc.settings.element.find(qn("w:" + tag))
        if old is not None:
            doc.settings.element.remove(old)
        original = settings.find(qn("w:" + tag))
        if original is not None:
            doc.settings.element.append(deepcopy(original))
    target_section = doc.sections[0]._sectPr
    for child in list(target_section):
        target_section.remove(child)
    for child in template_section:
        if child.tag in {qn("w:pgSz"), qn("w:pgMar"), qn("w:cols"), qn("w:docGrid")}:
            target_section.append(deepcopy(child))
    core = doc.core_properties
    core.title = content["title"]
    core.subject = "ASL PWI denoising comparison"
    core.author = content["authors"]
    core.last_modified_by = ""
    core.comments = "Typography follows the supplied Jinhan manuscript reference."
    core.created = datetime(2026, 10, 23, tzinfo=timezone.utc)
    core.modified = datetime(2026, 10, 24, tzinfo=timezone.utc)
    core.revision = 9

    for text, size in ((content["title"], 16), ("", 12)):
        p = paragraph(doc, text, size=size, bold=True, citations=False)
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.0
        pf.space_before = pf.space_after = Pt(15.6)
    for index, name in enumerate(content["authors"].split(", ")):
        if index:
            add_text(p, ", ", bold=True, citations=False)
        add_text(p, name, bold=True, citations=False)
        numbers = content.get("author_affiliations", {}).get(name, [])
        if numbers:
            run = set_font(p.add_run(",".join(map(str, numbers))), bold=True)
            run.font.superscript = True
    for number, affiliation in content.get("affiliations", {}).items():
        p = paragraph(doc)
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.line_spacing = 1.0
        pf.space_after = Pt(4)
        run = set_font(p.add_run(number), size=10.5)
        run.font.superscript = True
        add_text(p, affiliation, size=10.5, citations=False)
    heading(doc, "1 | Introduction", new_page=True)
    for text in content["introduction"]:
        paragraph(doc, text)
    heading(doc, "2 | Methods")
    for label, text in content["methods"]:
        paragraph(doc, text, label=label)
    heading(doc, "3 | Results")
    for text in content["results"]:
        paragraph(doc, text)
    heading(doc, "4 | Discussion and Conclusion")
    for text in content["discussion"]:
        discussion = paragraph(doc, text)
    heading(doc, "References", new_page=True)
    for index, text in enumerate(content["references"], 1):
        paragraph(doc, f"[{index}] {text}", citations=False)

    if content.get("code_availability_url"):
        heading(doc, "Acknowledgements")
        availability = paragraph(doc, "The code is available at ", citations=False)
        add_hyperlink(availability, content["code_availability_url"])
        add_text(availability, ".", citations=False)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    for index, (kind, rows) in enumerate(tables):
        add_table(doc, kind, rows, new_page=index > 0)
    for figure in figures:
        add_figure(doc, figure, args.figures_dir)

    apply_template_typography(doc, profiles, content)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        paragraphs += ["\t".join(c.text for c in row.cells) for row in table.rows]
    args.output.with_suffix(".txt").write_text("\n".join(paragraphs) + "\n", encoding="utf-8")
    print(f"Saved {args.output}: {len(doc.inline_shapes)} figures, {len(doc.tables)} tables")
    print(f"ISMRM 2027 text counts: {counts}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--content", type=Path, default=Path(__file__).with_name("manuscript_content.json"))
    parser.add_argument("--figures-manifest", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    build(parser.parse_args())


if __name__ == "__main__":
    main()
